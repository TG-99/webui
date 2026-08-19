import os
import urllib.request
import pickle
from contextlib import asynccontextmanager

import io
from fastapi import FastAPI, HTTPException, Depends, Query
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, date, timezone, timedelta
from bson import ObjectId
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database import get_db, init_db, verify_password, hash_password
from auth import create_access_token, get_current_user, require_admin



MODEL_FILES = [
    "tiny_face_detector_model-weights_manifest.json",
    "tiny_face_detector_model.bin",
    "face_landmark_68_model-weights_manifest.json",
    "face_landmark_68_model.bin",
    "face_recognition_model-weights_manifest.json",
    "face_recognition_model.bin",
    "ssd_mobilenetv1_model-weights_manifest.json",
    "ssd_mobilenetv1_model.bin"
]
MODEL_CDN_BASE = "https://cdn.jsdelivr.net/npm/@vladmandic/face-api@1.7.12/model/"
FACE_API_JS_CDN = "https://cdn.jsdelivr.net/npm/@vladmandic/face-api@1.7.12/dist/face-api.js"

def ensure_local_models():
    models_dir = os.path.join(os.path.dirname(__file__), "static", "models")
    os.makedirs(models_dir, exist_ok=True)

    # 1. Ensure local face-api.js library file is cached in static/models/
    js_lib_path = os.path.join(models_dir, "face-api.js")
    if not os.path.exists(js_lib_path) or os.path.getsize(js_lib_path) == 0:
        try:
            print("[Model Pre-loader] Caching face-api.js library to static/models/...")
            req = urllib.request.Request(FACE_API_JS_CDN, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
            with urllib.request.urlopen(req, timeout=20) as resp:
                data = resp.read()
                with open(js_lib_path, "wb") as f:
                    f.write(data)
        except Exception as e:
            print(f"[Model Pre-loader] Warning caching face-api.js: {e}")
    
    # 2. Ensure model weight files are cached
    for fname in MODEL_FILES:
        fpath = os.path.join(models_dir, fname)
        if not os.path.exists(fpath) or os.path.getsize(fpath) == 0:
            url = MODEL_CDN_BASE + fname
            try:
                print(f"[Model Pre-loader] Caching {fname} to local server...")
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
                with urllib.request.urlopen(req, timeout=15) as resp:
                    data = resp.read()
                    with open(fpath, "wb") as f:
                        f.write(data)
            except Exception as e:
                print(f"[Model Pre-loader] Warning downloading {fname}: {e}")

@asynccontextmanager
async def lifespan(app: FastAPI):
    try:
        init_db()
        print("[Startup] Database initialized successfully.")
    except Exception as e:
        print(f"[Startup Warning] Could not initialize database: {e}")

    uploads_path = os.path.join(os.path.dirname(__file__), "static", "uploads")
    os.makedirs(uploads_path, exist_ok=True)

    try:
        ensure_local_models()
        print("[Startup] AI Face recognition models verified and cached locally on server.")
    except Exception as e:
        print(f"[Startup Warning] Could not cache local models: {e}")

    yield


app = FastAPI(
    title="MOHAMMAD CONSTRUCTION & ENGINEERING - Worker Attendance System",
    version="1.0.0",
    lifespan=lifespan
)

# Enable CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# Helper to serialize Mongo object IDs
def serialize_doc(doc):
    if doc is None:
        return None
    doc["id"] = str(doc["_id"])
    del doc["_id"]
    return doc

# --- Pydantic Schemas ---
class LoginRequest(BaseModel):
    username: str
    password: str

class UserCreate(BaseModel):
    username: str
    password: str
    full_name: str
    role: str = "site_manager" # admin or site_manager
    phone: Optional[str] = ""
    email: Optional[str] = ""
    assigned_project_id: Optional[str] = ""
    assigned_project_name: Optional[str] = ""
    allowed_tabs: Optional[List[str]] = None

class UserUpdate(BaseModel):
    full_name: Optional[str] = None
    role: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    assigned_project_id: Optional[str] = None
    assigned_project_name: Optional[str] = None
    allowed_tabs: Optional[List[str]] = None
    password: Optional[str] = None

class ProjectModel(BaseModel):
    name: str
    location: Optional[str] = ""
    code: Optional[str] = ""
    status: Optional[str] = "Active"
    assigned_manager_id: Optional[str] = ""
    assigned_manager_name: Optional[str] = ""

class GroupModel(BaseModel):
    name: str
    description: Optional[str] = ""

class WorkerModel(BaseModel):
    name: str
    passport_number: Optional[str] = ""
    phone: Optional[str] = ""
    hourly_rate: Optional[float] = 0.0
    project_id: Optional[str] = ""
    group_id: Optional[str] = ""
    picture_url: Optional[str] = ""
    status: Optional[str] = "Active"

class AttendanceRecordRequest(BaseModel):
    worker_id: str
    date: str # YYYY-MM-DD
    check_in: Optional[str] = None # HH:MM or ISO timestamp
    check_out: Optional[str] = None # HH:MM or ISO timestamp
    status: Optional[str] = "Present" # Present, Late, Absent, Half-Day
    notes: Optional[str] = ""
    work_volume: Optional[str] = ""
    admin_confirmed: Optional[bool] = False

class BulkCheckInRequest(BaseModel):
    worker_ids: List[str]
    date: str # YYYY-MM-DD
    check_in_time: str # HH:MM
    admin_confirmed: Optional[bool] = False

class BulkCheckOutRequest(BaseModel):
    worker_ids: List[str]
    date: str # YYYY-MM-DD
    check_out_time: str # HH:MM
    admin_confirmed: Optional[bool] = False

class BulkAttendanceExportRequest(BaseModel):
    worker_ids: Optional[List[str]] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    project_name: Optional[str] = None
    group_name: Optional[str] = None
    title_banner: Optional[str] = None


class PhotoUploadRequest(BaseModel):
    worker_id: Optional[str] = ""
    image_data: str

# --- Auth Routes ---
@app.post("/api/auth/login")
def login(req: LoginRequest):
    db = get_db()
    user = db.users.find_one({"username": req.username})
    if not user or not verify_password(req.password, user["password_hash"]):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid username or password"
        )
    
    token = create_access_token({
        "sub": user["username"],
        "role": user.get("role", "site_manager")
    })
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": str(user["_id"]),
            "username": user["username"],
            "full_name": user.get("full_name", ""),
            "role": user.get("role", "site_manager"),
            "email": user.get("email", ""),
            "phone": user.get("phone", ""),
            "assigned_project_id": user.get("assigned_project_id", ""),
            "assigned_project_name": user.get("assigned_project_name", ""),
            "allowed_tabs": user.get("allowed_tabs", ["dashboard", "face-scanner", "attendance", "workers", "projects", "groups"])
        }
    }

@app.get("/api/auth/me")
def get_me(current_user: dict = Depends(get_current_user)):
    return current_user

def check_site_manager_access(current_user: dict, target_project_id: Optional[str] = None, target_worker_id: Optional[str] = None):
    if current_user.get("role") == "admin":
        return
    
    assigned_pid = current_user.get("assigned_project_id")
    if not assigned_pid:
        return
        
    db = get_db()
    if target_project_id and target_project_id != assigned_pid:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Forbidden: You are only allowed to manage data for your assigned construction site project."
        )
        
    if target_worker_id:
        try:
            w = db.workers.find_one({"_id": ObjectId(target_worker_id)})
            if w and w.get("project_id") and w["project_id"] != assigned_pid:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Forbidden: You are only allowed to manage workers at your assigned site project."
                )
        except HTTPException:
            raise
        except Exception:
            pass

# --- Projects Routes ---
@app.get("/api/projects")
def list_projects(current_user: dict = Depends(get_current_user)):
    db = get_db()
    projects = list(db.projects.find())
    return [serialize_doc(p) for p in projects]

@app.post("/api/projects")
def create_project(proj: ProjectModel, current_user: dict = Depends(get_current_user)):
    db = get_db()
    if db.projects.find_one({"name": proj.name}):
        raise HTTPException(status_code=400, detail="Project with this name already exists")
    
    doc = proj.dict()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = db.projects.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    if "_id" in doc: del doc["_id"]
    return doc

@app.put("/api/projects/{proj_id}")
def update_project(proj_id: str, proj: ProjectModel, current_user: dict = Depends(get_current_user)):
    db = get_db()
    try:
        oid = ObjectId(proj_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Project ID format")
    
    res = db.projects.update_one({"_id": oid}, {"$set": proj.dict()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Sync project_name on linked workers
    db.workers.update_many({"project_id": proj_id}, {"$set": {"project_name": proj.name}})
    
    updated = db.projects.find_one({"_id": oid})
    return serialize_doc(updated)

@app.delete("/api/projects/{proj_id}")
def delete_project(proj_id: str, current_user: dict = Depends(require_admin)):
    db = get_db()
    try:
        oid = ObjectId(proj_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Project ID format")
    
    db.projects.delete_one({"_id": oid})
    return {"message": "Project deleted successfully"}

# --- Groups Routes ---
@app.get("/api/groups")
def list_groups(current_user: dict = Depends(get_current_user)):
    db = get_db()
    groups = list(db.groups.find())
    return [serialize_doc(g) for g in groups]

@app.post("/api/groups")
def create_group(grp: GroupModel, current_user: dict = Depends(get_current_user)):
    db = get_db()
    if db.groups.find_one({"name": grp.name}):
        raise HTTPException(status_code=400, detail="Group with this name already exists")
    
    doc = grp.dict()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = db.groups.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    if "_id" in doc: del doc["_id"]
    return doc

@app.put("/api/groups/{grp_id}")
def update_group(grp_id: str, grp: GroupModel, current_user: dict = Depends(get_current_user)):
    db = get_db()
    try:
        oid = ObjectId(grp_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Group ID format")
    
    res = db.groups.update_one({"_id": oid}, {"$set": grp.dict()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Group not found")
    
    # Sync group_name on linked workers
    db.workers.update_many({"group_id": grp_id}, {"$set": {"group_name": grp.name}})
    
    updated = db.groups.find_one({"_id": oid})
    return serialize_doc(updated)

@app.delete("/api/groups/{grp_id}")
def delete_group(grp_id: str, current_user: dict = Depends(require_admin)):
    db = get_db()
    try:
        oid = ObjectId(grp_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Group ID format")
    
    db.groups.delete_one({"_id": oid})
    return {"message": "Worker Group deleted successfully"}

# --- Workers Routes ---
@app.get("/api/workers")
def list_workers(
    project_id: Optional[str] = None,
    group_id: Optional[str] = None,
    status: Optional[str] = None,
    query: Optional[str] = None,
    photo_status: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    if current_user.get("role") == "site_manager" and current_user.get("assigned_project_id"):
        project_id = current_user["assigned_project_id"]

    filter_query = {}
    if project_id:
        filter_query["project_id"] = project_id
    if group_id:
        filter_query["group_id"] = group_id
    if status:
        filter_query["status"] = status
    if query:
        filter_query["$or"] = [
            {"name": {"$regex": query, "$options": "i"}},
            {"passport_number": {"$regex": query, "$options": "i"}}
        ]
    if photo_status == "with_photo":
        filter_query["picture_url"] = {"$exists": True, "$ne": None, "$regex": "\\S+"}
    elif photo_status == "without_photo":
        filter_query["$or"] = [
            {"picture_url": {"$exists": False}},
            {"picture_url": None},
            {"picture_url": ""},
            {"picture_url": {"$regex": "^\\s*$"}}
        ]
    
    workers = list(db.workers.find(filter_query))
    return [serialize_doc(w) for w in workers]

@app.post("/api/workers")
def create_worker(worker: WorkerModel, current_user: dict = Depends(get_current_user)):
    db = get_db()
    if current_user.get("role") == "site_manager" and current_user.get("assigned_project_id"):
        worker.project_id = current_user["assigned_project_id"]

    passport_clean = (worker.passport_number or "").strip().upper()
    if passport_clean and db.workers.find_one({"passport_number": passport_clean}):
        raise HTTPException(status_code=400, detail="Worker with this Passport Number already exists")
    
    doc = worker.dict()
    doc["passport_number"] = passport_clean
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    
    # Resolve names if missing
    if doc.get("project_id"):
        try:
            p = db.projects.find_one({"_id": ObjectId(doc["project_id"])})
            if p: doc["project_name"] = p["name"]
        except: pass
    if doc.get("group_id"):
        try:
            g = db.groups.find_one({"_id": ObjectId(doc["group_id"])})
            if g: doc["group_name"] = g["name"]
        except: pass
        
    res = db.workers.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    if "_id" in doc: del doc["_id"]
    return doc

@app.get("/api/workers/{worker_id}")
def get_worker(worker_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    try:
        oid = ObjectId(worker_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Worker ID format")
    w = db.workers.find_one({"_id": oid})
    if not w:
        raise HTTPException(status_code=404, detail="Worker not found")
    return serialize_doc(w)

@app.put("/api/workers/{worker_id}")
def update_worker(worker_id: str, worker: WorkerModel, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_site_manager_access(current_user, target_worker_id=worker_id)
    
    if current_user.get("role") == "site_manager" and current_user.get("assigned_project_id"):
        worker.project_id = current_user["assigned_project_id"]

    try:
        oid = ObjectId(worker_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Worker ID format")
    
    passport_clean = (worker.passport_number or "").strip().upper()
    if passport_clean:
        existing = db.workers.find_one({"passport_number": passport_clean, "_id": {"$ne": oid}})
        if existing:
            raise HTTPException(status_code=400, detail="Worker with this Passport Number already exists")

    doc = worker.dict()
    doc["passport_number"] = passport_clean
    
    # Resolve names
    if doc.get("project_id"):
        try:
            p = db.projects.find_one({"_id": ObjectId(doc["project_id"])})
            if p: doc["project_name"] = p["name"]
        except: pass
    if doc.get("group_id"):
        try:
            g = db.groups.find_one({"_id": ObjectId(doc["group_id"])})
            if g: doc["group_name"] = g["name"]
        except: pass
        
    existing_worker = db.workers.find_one({"_id": oid})
    if not existing_worker:
        raise HTTPException(status_code=404, detail="Worker not found")

    old_pic = (existing_worker.get("picture_url") or "").strip()
    new_pic = (doc.get("picture_url") or "").strip()

    db.workers.update_one({"_id": oid}, {"$set": doc})
    
    # If picture_url was modified or removed, unset face_descriptor so AI engine re-indexes it
    if old_pic != new_pic:
        db.workers.update_one({"_id": oid}, {"$unset": {"face_descriptor": ""}})
        try:
            pickle_index = load_pickle_index()
            if worker_id in pickle_index:
                del pickle_index[worker_id]
                save_pickle_index(pickle_index)
        except Exception:
            pass

    updated = db.workers.find_one({"_id": oid})
    return serialize_doc(updated)

@app.delete("/api/workers/{worker_id}")
def delete_worker(worker_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_site_manager_access(current_user, target_worker_id=worker_id)
    try:
        oid = ObjectId(worker_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Worker ID format")
    
    db.workers.delete_one({"_id": oid})
    try:
        pickle_index = load_pickle_index()
        if worker_id in pickle_index:
            del pickle_index[worker_id]
            save_pickle_index(pickle_index)
    except Exception:
        pass
    return {"message": "Worker deleted successfully"}

@app.post("/api/workers/upload-photo")
def upload_worker_photo(req: PhotoUploadRequest, current_user: dict = Depends(get_current_user)):
    if not req.image_data:
        raise HTTPException(status_code=400, detail="Image data is required")
        
    try:
        img_data = req.image_data
        if "," in img_data:
            img_data = img_data.split(",")[1]
            
        binary_data = base64.b64decode(img_data)
        filename = f"worker_{uuid.uuid4().hex[:10]}.jpg"
        uploads_dir = os.path.join(os.path.dirname(__file__), "static", "uploads")
        os.makedirs(uploads_dir, exist_ok=True)
        
        filepath = os.path.join(uploads_dir, filename)
        with open(filepath, "wb") as f:
            f.write(binary_data)
            
        picture_url = f"/uploads/{filename}"
        
        if req.worker_id:
            try:
                db = get_db()
                db.workers.update_one({"_id": ObjectId(req.worker_id)}, {"$set": {"picture_url": picture_url}, "$unset": {"face_descriptor": ""}})
                pickle_index = load_pickle_index()
                if req.worker_id in pickle_index:
                    del pickle_index[req.worker_id]
                    save_pickle_index(pickle_index)
            except Exception:
                pass

        return {
            "success": True,
            "picture_url": picture_url,
            "message": "Worker photo uploaded successfully"
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to process photo upload: {str(e)}")

@app.get("/api/proxy-image")
def proxy_image(url: str = Query(...)):
    if not url or not url.strip():
        raise HTTPException(status_code=400, detail="URL parameter is required")
    try:
        req = urllib.request.Request(url.strip(), headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            content_type = resp.headers.get("Content-Type", "image/jpeg")
            data = resp.read()
            return Response(content=data, media_type=content_type, headers={
                "Cache-Control": "public, max-age=86400, immutable",
                "Access-Control-Allow-Origin": "*"
            })
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to proxy image: {str(e)}")

# ── Index.pickle Local Cache Helpers & Endpoints ──────────────────────────────
INDEX_PICKLE_PATH = os.path.join(os.path.dirname(__file__), "index.pickle")

def load_pickle_index() -> dict:
    if os.path.exists(INDEX_PICKLE_PATH):
        try:
            with open(INDEX_PICKLE_PATH, "rb") as f:
                data = pickle.load(f)
                if isinstance(data, dict):
                    return data
        except Exception as e:
            print(f"[Pickle Index] Warning loading index.pickle: {e}")
    return {}

def save_pickle_index(index_data: dict):
    try:
        with open(INDEX_PICKLE_PATH, "wb") as f:
            pickle.dump(index_data, f)
    except Exception as e:
        print(f"[Pickle Index] Error saving index.pickle: {e}")

class DescriptorCacheRequest(BaseModel):
    worker_id: str
    descriptor: List[float]

@app.get("/api/face-descriptors")
def get_face_descriptors(current_user: dict = Depends(get_current_user)):
    try:
        db = get_db()
        workers = list(db.workers.find({"status": {"$ne": "Terminated"}}))
        pickle_index = load_pickle_index()
        dirty_pickle = False

        result = []
        for w in workers:
            w_id = str(w["_id"])
            w["id"] = w_id
            del w["_id"]

            pic_url = (w.get("picture_url") or "").strip()

            # If photo was removed or is empty, worker MUST NOT have an indexed face descriptor!
            if not pic_url:
                descriptor = None
                if w_id in pickle_index:
                    del pickle_index[w_id]
                    dirty_pickle = True
                if "face_descriptor" in w:
                    db.workers.update_one({"_id": ObjectId(w_id)}, {"$unset": {"face_descriptor": ""}})
            else:
                descriptor = pickle_index.get(w_id)
                if not descriptor and "face_descriptor" in w and w["face_descriptor"]:
                    descriptor = w["face_descriptor"]
                    pickle_index[w_id] = descriptor
                    dirty_pickle = True

            result.append({
                "worker": w,
                "imgUrl": pic_url,
                "name": w.get("name", ""),
                "descriptor": descriptor
            })

        if dirty_pickle:
            save_pickle_index(pickle_index)

        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to load face descriptors: {str(e)}")

@app.post("/api/face-descriptors/cache")
def cache_face_descriptor(req: DescriptorCacheRequest, current_user: dict = Depends(get_current_user)):
    if not req.worker_id or not req.descriptor:
        raise HTTPException(status_code=400, detail="Worker ID and descriptor are required")
    try:
        pickle_index = load_pickle_index()
        pickle_index[req.worker_id] = req.descriptor
        save_pickle_index(pickle_index)

        db = get_db()
        db.workers.update_one({"_id": ObjectId(req.worker_id)}, {"$set": {"face_descriptor": req.descriptor}})
        return {"success": True, "message": "Face descriptor saved to index.pickle"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to cache face descriptor: {str(e)}")

@app.post("/api/face-descriptors/reset")
def reset_face_descriptors(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Only administrators can refresh the face index")
    try:
        # Reset index.pickle
        save_pickle_index({})
        
        # Unset face_descriptor on all worker documents in database
        db = get_db()
        db.workers.update_many({}, {"$unset": {"face_descriptor": ""}})
        
        return {"success": True, "message": "Face descriptors and index.pickle reset successfully."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to reset face descriptors: {str(e)}")



def check_24h_window(att_record: Optional[dict], target_date_str: str) -> tuple:
    """
    Determines if an attendance record or date is older than 24 hours from input recording.
    Returns (is_older_than_24h: bool, recorded_at_iso: Optional[str])
    """
    now_utc = datetime.now(timezone.utc)
    
    if att_record:
        rec_str = att_record.get("recorded_at") or att_record.get("created_at") or att_record.get("updated_at")
        if rec_str:
            try:
                clean_str = str(rec_str).replace("Z", "+00:00")
                rec_dt = datetime.fromisoformat(clean_str)
                if rec_dt.tzinfo is None:
                    rec_dt = rec_dt.replace(tzinfo=timezone.utc)
                elapsed_seconds = (now_utc - rec_dt).total_seconds()
                return (elapsed_seconds > 24 * 3600), rec_dt.isoformat()
            except Exception:
                pass

    # If no recorded_at timestamp exists yet, check against target_date midnight UTC + 36h
    try:
        t_date = datetime.strptime(target_date_str, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        elapsed_seconds = (now_utc - t_date).total_seconds()
        is_older = (elapsed_seconds > 36 * 3600)
        return is_older, None
    except Exception:
        return False, None

# --- Attendance Routes ---
@app.get("/api/attendance")
def get_attendance(
    target_date: Optional[str] = Query(None, alias="date"),
    project_id: Optional[str] = None,
    group_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    today_str = date.today().isoformat()
    if not target_date or target_date > today_str:
        target_date = today_str
    
    if current_user.get("role") == "site_manager" and current_user.get("assigned_project_id"):
        project_id = current_user["assigned_project_id"]

    # Fetch active workers filtered by project/group (Inactive workers like On Leave/Resigned do not appear on Daily Attendance)
    worker_filter = {
        "$or": [
            {"status": "Active"},
            {"status": {"$exists": False}}
        ]
    }
    if project_id: worker_filter["project_id"] = project_id
    if group_id: worker_filter["group_id"] = group_id
    
    workers = list(db.workers.find(worker_filter))
    worker_map = {str(w["_id"]): serialize_doc(w) for w in workers}
    worker_ids = list(worker_map.keys())
    
    # Fetch existing attendance records for the target date
    records = list(db.attendance.find({
        "date": target_date,
        "worker_id": {"$in": worker_ids}
    }))
    
    record_map = {r["worker_id"]: serialize_doc(r) for r in records}
    
    # Merge workers with their attendance for target_date
    result = []
    for wid, wdata in worker_map.items():
        att = record_map.get(wid)
        if att:
            is_older_than_24h, rec_at_iso = check_24h_window(att, target_date)
            att["recorded_at"] = att.get("recorded_at") or rec_at_iso
            att["is_older_than_24h"] = is_older_than_24h
            att["is_locked"] = is_older_than_24h and (current_user.get("role") == "site_manager")
            if not att.get("check_in"):
                att["status"] = "Absent"
            elif att.get("status") in [None, "", "Absent", "Pending"]:
                att["status"] = "Present"
        else:
            is_older_than_24h, rec_at_iso = check_24h_window(None, target_date)
            att = {
                "worker_id": wid,
                "date": target_date,
                "check_in": None,
                "check_out": None,
                "status": "Absent", # Absent if no check-in
                "hours_worked": 0.0,
                "notes": "",
                "work_volume": "",
                "recorded_at": None,
                "is_older_than_24h": is_older_than_24h,
                "is_locked": is_older_than_24h and (current_user.get("role") == "site_manager")
            }

        if att.get("check_in") and att.get("check_out"):
            att["hours_worked"] = calculate_hours(att["check_in"], att["check_out"])
            
        result.append({
            "worker": wdata,
            "attendance": att
        })
        
    return {
        "date": target_date,
        "records": result
    }

def get_16th_cycle_start_date(ref_date: Optional[date] = None) -> str:
    if ref_date is None:
        ref_date = date.today()
    if ref_date.day >= 16:
        start_d = ref_date.replace(day=16)
    else:
        month = ref_date.month - 1
        year = ref_date.year
        if month < 1:
            month = 12
            year -= 1
        start_d = date(year, month, 16)
    return start_d.isoformat()

@app.get("/api/attendance/worker/{worker_id}")
def get_worker_attendance_history(
    worker_id: str,
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    check_site_manager_access(current_user, target_worker_id=worker_id)

    try:
        oid = ObjectId(worker_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid Worker ID format")

    worker = db.workers.find_one({"_id": oid})
    if not worker:
        raise HTTPException(status_code=404, detail="Worker not found")
    worker_doc = serialize_doc(worker)

    today_str = date.today().isoformat()

    if end_date and isinstance(end_date, str):
        eff_end = min(end_date, today_str)
    else:
        eff_end = today_str

    records_raw = list(db.attendance.find({"worker_id": worker_id}))

    if start_date and isinstance(start_date, str):
        eff_start = start_date
    else:
        default_start = get_16th_cycle_start_date()
        dates_found = []
        for r in records_raw:
            if r.get("date"):
                dates_found.append(r["date"])
        if worker.get("created_at"):
            try:
                c_date = str(worker["created_at"])[:10]
                if len(c_date) == 10 and c_date.count("-") == 2:
                    dates_found.append(c_date)
            except:
                pass
        if dates_found:
            min_found = min(dates_found)
            eff_start = min(min_found, default_start)
        else:
            eff_start = default_start

    if eff_start > eff_end:
        eff_start = eff_end

    record_map = {}
    for r in records_raw:
        doc = serialize_doc(r)
        d = doc.get("date")
        if d and eff_start <= d <= eff_end:
            record_map[d] = doc

    try:
        start_dt = datetime.strptime(eff_start, "%Y-%m-%d").date()
        end_dt = datetime.strptime(eff_end, "%Y-%m-%d").date()
    except Exception:
        start_dt = datetime.strptime(get_16th_cycle_start_date(), "%Y-%m-%d").date()
        end_dt = date.today()

    records = []
    curr_dt = start_dt
    while curr_dt <= end_dt:
        d_str = curr_dt.isoformat()
        if d_str in record_map:
            doc = record_map[d_str]
            if doc.get("check_in") and doc.get("check_out"):
                doc["hours_worked"] = calculate_hours(doc["check_in"], doc["check_out"])
            elif not doc.get("check_in"):
                doc["status"] = doc.get("status") or "Absent"
                doc["hours_worked"] = float(doc.get("hours_worked") or 0.0)
            records.append(doc)
        else:
            records.append({
                "worker_id": worker_id,
                "date": d_str,
                "check_in": None,
                "check_out": None,
                "status": "Absent",
                "hours_worked": 0.0,
                "notes": "",
                "work_volume": "",
                "recorded_at": None
            })
        curr_dt += timedelta(days=1)

    total_hours = round(sum(float(r.get("hours_worked") or 0) for r in records), 2)
    days_worked = sum(
        1 for r in records if r.get("status") in ["Present", "Late", "Half-Day"] or r.get("check_in")
    )

    return {
        "worker": worker_doc,
        "records": records,
        "total_hours": total_hours,
        "days_worked": days_worked
    }

def fetch_bulk_attendance_matrix_data(db, current_user, worker_ids=None, start_date=None, end_date=None, project_name=None, group_name=None):
    today_str = date.today().isoformat()
    if not start_date:
        start_date = get_16th_cycle_start_date()
    if not end_date:
        end_date = today_str

    if start_date > end_date:
        start_date = end_date

    # Build worker query
    query = {}
    if worker_ids and len(worker_ids) > 0:
        oids = []
        for wid in worker_ids:
            try: oids.append(ObjectId(wid))
            except: pass
        if oids:
            query["_id"] = {"$in": oids}
    
    if current_user.get("role") == "site_manager" and current_user.get("assigned_project_id"):
        try:
            query["project_id"] = current_user["assigned_project_id"]
        except: pass

    # Fetch workers in addition order (natural MongoDB insertion order)
    workers_cursor = db.workers.find(query)
    workers = [serialize_doc(w) for w in workers_cursor]

    # If specific worker_ids order was provided, preserve that exact requested order
    if worker_ids and len(worker_ids) > 0:
        order_map = {wid: idx for idx, wid in enumerate(worker_ids)}
        workers.sort(key=lambda w: order_map.get(w["id"], 999999))

    # Filter further by project_name / group_name if provided
    if project_name and project_name != "All":
        workers = [w for w in workers if w.get("project_name") == project_name]
    if group_name and group_name != "All":
        workers = [w for w in workers if w.get("group_name") == group_name]
    else:
        # Group workers by (project_name, group_name) maintaining addition order of groups
        grouped = {}
        for w in workers:
            key = ((w.get("project_name") or "").strip(), (w.get("group_name") or "").strip())
            if key not in grouped:
                grouped[key] = []
            grouped[key].append(w)
        ordered_workers = []
        for grp in grouped.values():
            ordered_workers.extend(grp)
        workers = ordered_workers

    # Generate dates range
    try:
        dt_start = datetime.strptime(start_date, "%Y-%m-%d").date()
        dt_end = datetime.strptime(end_date, "%Y-%m-%d").date()
    except Exception:
        dt_start = datetime.strptime(get_16th_cycle_start_date(), "%Y-%m-%d").date()
        dt_end = date.today()

    DAYS_SHORT = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    dates = []
    curr = dt_start
    while curr <= dt_end:
        iso_str = curr.isoformat()
        day_idx = curr.weekday() # 0 = Mon, 6 = Sun
        day_str = DAYS_SHORT[day_idx]
        fmt_str = curr.strftime("%d.%m.%Y")
        is_sun = (day_idx == 6)
        dates.append({
            "date": iso_str,
            "formatted": fmt_str,
            "day": day_str,
            "header": f"{fmt_str}, {day_str}",
            "is_sunday": is_sun
        })
        curr += timedelta(days=1)

    target_worker_ids = [w["id"] for w in workers if "id" in w]

    # Query attendance records for these workers and date range
    att_cursor = db.attendance.find({
        "worker_id": {"$in": target_worker_ids},
        "date": {"$gte": start_date, "$lte": end_date}
    })
    att_records = [serialize_doc(r) for r in att_cursor]

    # Map attendance: worker_id -> date -> hours_worked
    att_map = {}
    for r in att_records:
        wid = r.get("worker_id")
        d = r.get("date")
        if not wid or not d:
            continue
        if wid not in att_map:
            att_map[wid] = {}
        
        h = float(r.get("hours_worked") or 0.0)
        if h <= 0 and r.get("check_in") and r.get("check_out"):
            h = calculate_hours(r.get("check_in"), r.get("check_out"))
        att_map[wid][d] = h

    # Build matrix per worker
    matrix = {}
    grand_total_hours = 0.0
    for w in workers:
        wid = w["id"]
        w_hours_map = att_map.get(wid, {})
        daily_hours = {}
        row_total = 0.0
        for d_info in dates:
            d_str = d_info["date"]
            h = float(w_hours_map.get(d_str, 0.0))
            daily_hours[d_str] = round(h, 2)
            row_total += h
        
        row_total = round(row_total, 2)
        grand_total_hours += row_total
        matrix[wid] = {
            "daily_hours": daily_hours,
            "total_hours": row_total
        }

    p_title = project_name if (project_name and project_name != "All") else ""
    g_title = group_name if (group_name and group_name != "All") else ""
    if g_title and p_title:
        banner = f"{g_title} - {p_title}"
    elif g_title:
        banner = g_title
    elif p_title:
        banner = p_title
    else:
        banner = "All Groups & Projects"

    return {
        "start_date": start_date,
        "end_date": end_date,
        "banner": banner,
        "dates": dates,
        "workers": workers,
        "matrix": matrix,
        "grand_total_hours": round(grand_total_hours, 2)
    }

@app.get("/api/attendance/bulk-history")
def get_bulk_attendance_history(
    worker_ids: Optional[str] = Query(None),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    project_name: Optional[str] = Query(None),
    group_name: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    w_list = [w.strip() for w in worker_ids.split(",") if w.strip()] if worker_ids else None
    return fetch_bulk_attendance_matrix_data(db, current_user, w_list, start_date, end_date, project_name, group_name)

@app.post("/api/attendance/export-bulk-excel")
def export_bulk_attendance_excel(
    req: BulkAttendanceExportRequest,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    data = fetch_bulk_attendance_matrix_data(
        db, current_user,
        worker_ids=req.worker_ids,
        start_date=req.start_date,
        end_date=req.end_date,
        project_name=req.project_name,
        group_name=req.group_name
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Matrix"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
    font_subtitle = Font(name="Calibri", size=11, bold=True, color="595959")
    font_header_normal = Font(name="Calibri", size=9, bold=True, color="000000")
    font_header_sunday = Font(name="Calibri", size=9, bold=True, color="FF0000")
    font_banner = Font(name="Calibri", size=11, bold=True, color="000000")
    font_worker_name = Font(name="Calibri", size=10, bold=True, color="000000")
    font_data_normal = Font(name="Calibri", size=10, color="000000")
    font_data_sunday = Font(name="Calibri", size=10, color="FF0000")
    font_total = Font(name="Calibri", size=10, bold=True, color="000000")

    fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_worker_name = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_banner = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
    fill_total = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="BFBFBF")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A1:D1")
    ws["A1"] = "MOHAMMAD CONSTRUCTION & ENGINEERING SDN.BHD."
    ws["A1"].font = font_title

    ws.merge_cells("A2:D2")
    ws["A2"] = "WORKER ATTENDANCE HISTORY REPORT"
    ws["A2"].font = font_subtitle

    start_fmt = datetime.strptime(data["start_date"], "%Y-%m-%d").strftime("%d.%m.%Y")
    end_fmt = datetime.strptime(data["end_date"], "%Y-%m-%d").strftime("%d.%m.%Y")
    ws.merge_cells("A3:D3")
    ws["A3"] = f"Period: {start_fmt} to {end_fmt}"
    ws["A3"].font = font_subtitle

    row_header = 5
    ws.cell(row=row_header, column=1, value="NO.").alignment = align_center
    ws.cell(row=row_header, column=2, value="NAME").alignment = align_left
    ws.cell(row=row_header, column=3, value="PASSPORT No.").alignment = align_center

    dates = data["dates"]
    for i, d_info in enumerate(dates):
        c_idx = 4 + i
        cell = ws.cell(row=row_header, column=c_idx, value=d_info["header"])
        cell.alignment = align_center
        cell.font = font_header_sunday if d_info["is_sunday"] else font_header_normal
        cell.fill = fill_header
        cell.border = thin_border

    total_col_idx = 4 + len(dates)
    cell_tot_h = ws.cell(row=row_header, column=total_col_idx, value="TOTAL HOURS")
    cell_tot_h.alignment = align_center
    cell_tot_h.font = font_total
    cell_tot_h.fill = fill_header
    cell_tot_h.border = thin_border

    for col in range(1, 4):
        c = ws.cell(row=row_header, column=col)
        c.font = font_header_normal
        c.fill = fill_header
        c.border = thin_border

    last_col_letter = get_column_letter(total_col_idx)
    workers = data["workers"]
    matrix = data["matrix"]
    current_row = 6

    is_all_groups = not req.group_name or req.group_name == "All"

    if is_all_groups:
        group_map = {}
        for w in workers:
            p_name = (w.get("project_name") or "").strip()
            g_name = (w.get("group_name") or "").strip()
            if p_name and g_name:
                label = f"{p_name} - {g_name}"
            elif g_name:
                label = g_name
            elif p_name:
                label = p_name
            else:
                label = "Unassigned"

            if label not in group_map:
                group_map[label] = []
            group_map[label].append(w)

        for group_label, group_workers in group_map.items():
            ws.merge_cells(f"A{current_row}:{last_col_letter}{current_row}")
            grp_cell = ws.cell(row=current_row, column=1, value=group_label.upper())
            grp_cell.font = font_banner
            grp_cell.alignment = align_center
            grp_cell.fill = fill_banner
            for c in range(1, total_col_idx + 1):
                ws.cell(row=current_row, column=c).border = thin_border
            current_row += 1

            for idx, w in enumerate(group_workers):
                wid = w["id"]
                w_data = matrix.get(wid, {"daily_hours": {}, "total_hours": 0.0})
                daily_hours = w_data["daily_hours"]
                tot = w_data["total_hours"]

                c1 = ws.cell(row=current_row, column=1, value=idx + 1)
                c1.alignment = align_center
                c1.font = font_data_normal
                c1.border = thin_border

                c2 = ws.cell(row=current_row, column=2, value=str(w.get("name", "")).upper())
                c2.alignment = align_left
                c2.font = font_worker_name
                c2.fill = fill_worker_name
                c2.border = thin_border

                c3 = ws.cell(row=current_row, column=3, value=w.get("passport_number") or "—")
                c3.alignment = align_center
                c3.font = font_data_normal
                c3.border = thin_border

                for i, d_info in enumerate(dates):
                    c_idx = 4 + i
                    d_str = d_info["date"]
                    val = daily_hours.get(d_str, 0.0)
                    val_display = int(val) if isinstance(val, (int, float)) and float(val).is_integer() else val
                    cell = ws.cell(row=current_row, column=c_idx, value=val_display)
                    cell.alignment = align_center
                    cell.font = font_data_sunday if d_info["is_sunday"] else font_data_normal
                    cell.border = thin_border

                tot_display = int(tot) if isinstance(tot, (int, float)) and float(tot).is_integer() else tot
                c_tot = ws.cell(row=current_row, column=total_col_idx, value=tot_display)
                c_tot.alignment = align_center
                c_tot.font = font_total
                c_tot.fill = fill_total
                c_tot.border = thin_border

                current_row += 1
    else:
        for idx, w in enumerate(workers):
            wid = w["id"]
            w_data = matrix.get(wid, {"daily_hours": {}, "total_hours": 0.0})
            daily_hours = w_data["daily_hours"]
            tot = w_data["total_hours"]

            c1 = ws.cell(row=current_row, column=1, value=idx + 1)
            c1.alignment = align_center
            c1.font = font_data_normal
            c1.border = thin_border

            c2 = ws.cell(row=current_row, column=2, value=str(w.get("name", "")).upper())
            c2.alignment = align_left
            c2.font = font_worker_name
            c2.fill = fill_worker_name
            c2.border = thin_border

            c3 = ws.cell(row=current_row, column=3, value=w.get("passport_number") or "—")
            c3.alignment = align_center
            c3.font = font_data_normal
            c3.border = thin_border

            for i, d_info in enumerate(dates):
                c_idx = 4 + i
                d_str = d_info["date"]
                val = daily_hours.get(d_str, 0.0)
                val_display = int(val) if isinstance(val, (int, float)) and float(val).is_integer() else val
                cell = ws.cell(row=current_row, column=c_idx, value=val_display)
                cell.alignment = align_center
                cell.font = font_data_sunday if d_info["is_sunday"] else font_data_normal
                cell.border = thin_border

            tot_display = int(tot) if isinstance(tot, (int, float)) and float(tot).is_integer() else tot
            c_tot = ws.cell(row=current_row, column=total_col_idx, value=tot_display)
            c_tot.alignment = align_center
            c_tot.font = font_total
            c_tot.fill = fill_total
            c_tot.border = thin_border

            current_row += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 16
    for i in range(len(dates)):
        col_let = get_column_letter(4 + i)
        ws.column_dimensions[col_let].width = 12
    ws.column_dimensions[last_col_letter].width = 14

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Worker_Attendance_Matrix_{data['start_date']}_to_{data['end_date']}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def calculate_hours(check_in_str, check_out_str):
    if not check_in_str or not check_out_str:
        return 0.0
    
    t1 = None
    t2 = None
    
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I:%M:%S %p"):
        if not t1:
            try: t1 = datetime.strptime(str(check_in_str).strip(), fmt)
            except ValueError: pass
        if not t2:
            try: t2 = datetime.strptime(str(check_out_str).strip(), fmt)
            except ValueError: pass
        if t1 and t2:
            break
            
    if not t1 or not t2:
        return 0.0
        
    delta_sec = (t2 - t1).total_seconds()
    if delta_sec < 0:
        delta_sec += 86400  # Overnight shift
        
    hours = delta_sec / 3600.0
    return round(max(0.0, hours), 2)

@app.post("/api/attendance/record")
def record_attendance(req: AttendanceRecordRequest, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_site_manager_access(current_user, target_worker_id=req.worker_id)
    
    today_str = date.today().isoformat()
    if req.date > today_str:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Future dates are disabled for attendance recording."
        )

    check_in_clean = (req.check_in or "").strip() if req.check_in else None
    check_out_clean = (req.check_out or "").strip() if req.check_out else None

    if check_in_clean and check_out_clean:
        if check_in_clean == check_out_clean:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Invalid attendance times: Check-In time ({check_in_clean}) and Check-Out time ({check_out_clean}) cannot be identical."
            )

    if check_in_clean:
        try:
            worker_obj = db.workers.find_one({"_id": ObjectId(req.worker_id)})
            if worker_obj and worker_obj.get("status") in ["On Leave", "Resigned", "Terminated"]:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Worker '{worker_obj.get('name', 'Worker')}' is currently {worker_obj.get('status')} and cannot be checked in."
                )
        except HTTPException:
            raise
        except Exception:
            pass

    filter_doc = {"worker_id": req.worker_id, "date": req.date}
    existing = db.attendance.find_one(filter_doc)

    is_older_than_24h, recorded_at_str = check_24h_window(existing, req.date)

    if is_older_than_24h:
        if current_user.get("role") == "site_manager":
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Editing locked: Site managers can only edit worker work details within 24 hours of input recording."
            )
        elif current_user.get("role") == "admin":
            if not req.admin_confirmed:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="24_HOUR_OVERRIDE_REQUIRED: Work details were recorded over 24 hours ago. Admin confirmation is required."
                )

    now_iso = datetime.now(timezone.utc).isoformat()
    if existing:
        recorded_at = existing.get("recorded_at") or recorded_at_str or existing.get("created_at") or existing.get("updated_at") or now_iso
    else:
        recorded_at = now_iso

    # If both Check-In and Check-Out times are empty/null, delete attendance record from database
    if not check_in_clean and not check_out_clean:
        if existing:
            db.attendance.delete_one(filter_doc)
        return {
            "worker_id": req.worker_id,
            "date": req.date,
            "check_in": None,
            "check_out": None,
            "status": "Pending",
            "hours_worked": 0.0,
            "notes": "",
            "deleted": True,
            "recorded_at": recorded_at,
            "updated_at": None,
            "is_older_than_24h": is_older_than_24h,
            "is_locked": False
        }

    hours = calculate_hours(check_in_clean, check_out_clean)
    if check_in_clean:
        status_val = req.status if (req.status and req.status not in ["Pending", "Absent"]) else ("Late" if check_in_clean > "09:00" else "Present")
    else:
        status_val = "Absent"
    
    update_doc = {
        "$set": {
            "worker_id": req.worker_id,
            "date": req.date,
            "check_in": check_in_clean,
            "check_out": check_out_clean,
            "status": status_val,
            "hours_worked": hours,
            "notes": req.notes,
            "work_volume": req.work_volume or "",
            "recorded_at": recorded_at,
            "updated_by": current_user["username"],
            "updated_at": now_iso
        }
    }
    
    db.attendance.update_one(filter_doc, update_doc, upsert=True)
    saved = db.attendance.find_one(filter_doc)
    doc = serialize_doc(saved)
    doc["is_older_than_24h"] = is_older_than_24h
    doc["is_locked"] = is_older_than_24h and (current_user.get("role") == "site_manager")
    return doc

@app.post("/api/attendance/bulk-checkin")
def bulk_checkin(req: BulkCheckInRequest, current_user: dict = Depends(get_current_user)):
    db = get_db()
    
    today_str = date.today().isoformat()
    if req.date > today_str:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Future dates are disabled for attendance recording."
        )

    # Check if any workers have existing records > 24 hours old
    locked_records_exist = False
    for wid in req.worker_ids:
        check_site_manager_access(current_user, target_worker_id=wid)
        existing = db.attendance.find_one({"worker_id": wid, "date": req.date})
        is_older, _ = check_24h_window(existing, req.date)
        if is_older:
            locked_records_exist = True
            break

    if locked_records_exist:
        if current_user.get("role") == "site_manager":
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Editing locked: Attendance records for one or more workers on this date were recorded >24 hours ago."
            )
        elif current_user.get("role") == "admin":
            if not req.admin_confirmed:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="24_HOUR_OVERRIDE_REQUIRED: One or more attendance records were recorded over 24 hours ago. Admin confirmation is required."
                )

    # Filter active workers (exclude On Leave / Resigned / Terminated workers)
    active_worker_ids = set()
    try:
        w_oids = []
        for wid in req.worker_ids:
            try: w_oids.append(ObjectId(wid))
            except: pass
        active_docs = list(db.workers.find({
            "_id": {"$in": w_oids},
            "$or": [
                {"status": "Active"},
                {"status": {"$exists": False}}
            ]
        }))
        active_worker_ids = {str(w["_id"]) for w in active_docs}
    except Exception:
        pass

    now_iso = datetime.now(timezone.utc).isoformat()
    updated_count = 0
    
    for wid in req.worker_ids:
        if active_worker_ids and wid not in active_worker_ids:
            continue

        filter_doc = {"worker_id": wid, "date": req.date}
        existing = db.attendance.find_one(filter_doc)
        
        recorded_at = now_iso
        if existing:
            rec_str = existing.get("recorded_at") or existing.get("created_at") or existing.get("updated_at")
            if rec_str: recorded_at = rec_str

        check_out = existing.get("check_out") if existing else None
        if check_out and str(check_out).strip() == str(req.check_in_time).strip():
            check_out = None
        hours = calculate_hours(req.check_in_time, check_out)
        
        db.attendance.update_one(
            filter_doc,
            {"$set": {
                "worker_id": wid,
                "date": req.date,
                "check_in": req.check_in_time,
                "check_out": check_out,
                "status": "Present",
                "hours_worked": hours,
                "recorded_at": recorded_at,
                "updated_by": current_user["username"],
                "updated_at": now_iso
            }},
            upsert=True
        )
        updated_count += 1
        
    return {"message": f"Successfully checked in {updated_count} workers", "count": updated_count}

@app.post("/api/attendance/bulk-checkout")
def bulk_checkout(req: BulkCheckOutRequest, current_user: dict = Depends(get_current_user)):
    db = get_db()
    
    today_str = date.today().isoformat()
    if req.date > today_str:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Future dates are disabled for attendance recording."
        )

    # Check if any workers have existing records > 24 hours old
    locked_records_exist = False
    for wid in req.worker_ids:
        check_site_manager_access(current_user, target_worker_id=wid)
        existing = db.attendance.find_one({"worker_id": wid, "date": req.date})
        is_older, _ = check_24h_window(existing, req.date)
        if is_older:
            locked_records_exist = True
            break

    if locked_records_exist:
        if current_user.get("role") == "site_manager":
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Editing locked: Attendance records for one or more workers on this date were recorded >24 hours ago."
            )
        elif current_user.get("role") == "admin":
            if not req.admin_confirmed:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="24_HOUR_OVERRIDE_REQUIRED: One or more attendance records were recorded over 24 hours ago. Admin confirmation is required."
                )

    # Filter active workers (exclude On Leave / Resigned / Terminated workers)
    active_worker_ids = set()
    try:
        w_oids = []
        for wid in req.worker_ids:
            try: w_oids.append(ObjectId(wid))
            except: pass
        active_docs = list(db.workers.find({
            "_id": {"$in": w_oids},
            "$or": [
                {"status": "Active"},
                {"status": {"$exists": False}}
            ]
        }))
        active_worker_ids = {str(w["_id"]) for w in active_docs}
    except Exception:
        pass

    now_iso = datetime.now(timezone.utc).isoformat()
    updated_count = 0
    
    for wid in req.worker_ids:
        if active_worker_ids and wid not in active_worker_ids:
            continue

        filter_doc = {"worker_id": wid, "date": req.date}
        existing = db.attendance.find_one(filter_doc)
        
        # Batch Check-Out applies ONLY to workers who have checked in! Skip workers with no check-in.
        check_in = existing.get("check_in") if existing else None
        if not check_in or not str(check_in).strip():
            continue

        # Skip workers where check_in is identical to batch check_out_time
        if str(check_in).strip() == str(req.check_out_time).strip():
            continue

        recorded_at = now_iso
        if existing:
            rec_str = existing.get("recorded_at") or existing.get("created_at") or existing.get("updated_at")
            if rec_str: recorded_at = rec_str

        hours = calculate_hours(check_in, req.check_out_time)
        
        db.attendance.update_one(
            filter_doc,
            {"$set": {
                "worker_id": wid,
                "date": req.date,
                "check_in": check_in,
                "check_out": req.check_out_time,
                "status": "Present",
                "hours_worked": hours,
                "recorded_at": recorded_at,
                "updated_by": current_user["username"],
                "updated_at": now_iso
            }},
            upsert=True
        )
        updated_count += 1
        
    return {"message": f"Successfully checked out {updated_count} workers", "count": updated_count}

@app.get("/api/attendance/stats")
def get_stats(target_date: Optional[str] = Query(None, alias="date"), current_user: dict = Depends(get_current_user)):
    db = get_db()
    today_str = date.today().isoformat()
    if not target_date or target_date > today_str:
        target_date = today_str
        
    active_filter = {"$or": [{"status": "Active"}, {"status": {"$exists": False}}]}
    if current_user.get("role") == "site_manager" and current_user.get("assigned_project_id"):
        assigned_pid = current_user["assigned_project_id"]
        total_workers = db.workers.count_documents({"project_id": assigned_pid, **active_filter})
        site_worker_ids = [str(w["_id"]) for w in db.workers.find({"project_id": assigned_pid, **active_filter})]
        records = list(db.attendance.find({"date": target_date, "worker_id": {"$in": site_worker_ids}}))
    else:
        total_workers = db.workers.count_documents(active_filter)
        records = list(db.attendance.find({"date": target_date}))
    
    present = sum(1 for r in records if r.get("check_in"))
    late = sum(1 for r in records if r.get("status") == "Late" or (r.get("check_in") and r.get("check_in") > "09:00"))
    absent = max(0, total_workers - present)
    checked_out = sum(1 for r in records if r.get("check_out"))
    pending = 0
    
    return {
        "date": target_date,
        "total_workers": total_workers,
        "present": present,
        "late": late,
        "absent": absent,
        "checked_out": checked_out,
        "pending": pending
    }

# --- Admin User Management Routes ---
@app.get("/api/admin/users")
def list_users(current_user: dict = Depends(require_admin)):
    db = get_db()
    users = list(db.users.find())
    for u in users:
        u["id"] = str(u["_id"])
        del u["_id"]
        if "password_hash" in u: del u["password_hash"]
    return users

@app.post("/api/admin/users")
def create_user(u: UserCreate, current_user: dict = Depends(require_admin)):
    db = get_db()
    uname = u.username.strip().lower()
    if db.users.find_one({"username": uname}):
        raise HTTPException(status_code=400, detail="Username already exists")
    
    proj_name = u.assigned_project_name
    if u.assigned_project_id:
        try:
            p = db.projects.find_one({"_id": ObjectId(u.assigned_project_id)})
            if p: proj_name = p["name"]
        except: pass
        
    user_doc = {
        "username": uname,
        "password_hash": hash_password(u.password),
        "full_name": u.full_name,
        "role": u.role,
        "phone": u.phone or u.email or "",
        "email": u.email or "",
        "assigned_project_id": u.assigned_project_id or "",
        "assigned_project_name": proj_name or "",
        "allowed_tabs": u.allowed_tabs if u.allowed_tabs is not None else ["dashboard", "attendance", "workers", "projects", "groups"],
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    res = db.users.insert_one(user_doc)
    user_doc["id"] = str(res.inserted_id)
    del user_doc["_id"]
    del user_doc["password_hash"]
    return user_doc

@app.put("/api/admin/users/{username}")
def update_user(username: str, u: UserUpdate, current_user: dict = Depends(require_admin)):
    db = get_db()
    existing = db.users.find_one({"username": username})
    if not existing:
        raise HTTPException(status_code=404, detail="User not found")
        
    if u.role is not None and u.role != existing.get("role") and existing.get("role") == "admin":
        admin_count = db.users.count_documents({"role": "admin"})
        if admin_count <= 1:
            raise HTTPException(
                status_code=400,
                detail="Cannot change role: System requires at least one active Admin account."
            )

    update_data = {}
    if u.full_name is not None: update_data["full_name"] = u.full_name
    if u.role is not None: update_data["role"] = u.role
    if u.phone is not None: update_data["phone"] = u.phone
    if u.email is not None: update_data["email"] = u.email
    if u.allowed_tabs is not None: update_data["allowed_tabs"] = u.allowed_tabs
    
    if u.assigned_project_id is not None:
        update_data["assigned_project_id"] = u.assigned_project_id
        if u.assigned_project_id:
            try:
                p = db.projects.find_one({"_id": ObjectId(u.assigned_project_id)})
                if p: update_data["assigned_project_name"] = p["name"]
            except: pass
        else:
            update_data["assigned_project_name"] = ""
            
    if u.password and u.password.strip():
        update_data["password_hash"] = hash_password(u.password.strip())
        
    db.users.update_one({"username": username}, {"$set": update_data})
    updated = db.users.find_one({"username": username})
    updated["id"] = str(updated["_id"])
    del updated["_id"]
    if "password_hash" in updated: del updated["password_hash"]
    return updated

@app.delete("/api/admin/users/{username}")
def delete_user(username: str, current_user: dict = Depends(require_admin)):
    db = get_db()
    if username == current_user["username"]:
        raise HTTPException(status_code=400, detail="Cannot delete your own admin account")
    
    target_user = db.users.find_one({"username": username})
    if target_user and target_user.get("role") == "admin":
        admin_count = db.users.count_documents({"role": "admin"})
        if admin_count <= 1:
            raise HTTPException(status_code=400, detail="Cannot delete the sole Admin account in the system.")

    res = db.users.delete_one({"username": username})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    return {"message": f"User {username} removed successfully"}

# --- Static Files & AI Models Mount ---
class CachedStaticFiles(StaticFiles):
    async def get_response(self, path: str, scope):
        response = await super().get_response(path, scope)
        response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
        response.headers["Access-Control-Allow-Origin"] = "*"
        response.headers["Content-Disposition"] = "inline"
        if path.endswith(".bin"):
            response.headers["Content-Type"] = "application/octet-stream"
        elif path.endswith(".json"):
            response.headers["Content-Type"] = "application/json"
        elif path.endswith(".js"):
            response.headers["Content-Type"] = "application/javascript"
        return response

models_dir_path = os.path.join(os.path.dirname(__file__), "static", "models")
if not os.path.exists(models_dir_path):
    os.makedirs(models_dir_path, exist_ok=True)

app.mount("/models", CachedStaticFiles(directory=models_dir_path), name="models")

static_dir = os.path.join(os.path.dirname(__file__), "static")
if not os.path.exists(static_dir):
    os.makedirs(static_dir, exist_ok=True)

app.mount("/", StaticFiles(directory=static_dir, html=True), name="static")
